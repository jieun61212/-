from io import BytesIO
from docx import Document
from app.models import Article


def txt_bytes(article, demo=False):
    label = '[예시 모드 · AI 생성 및 법률 사실 검증 미실행]\n\n' if demo else ''
    return (label + article.text).encode('utf-8-sig')


def docx_bytes(article: Article, demo=False):
    doc = Document()
    from docx.shared import Pt
    from docx.oxml.ns import qn
    normal = doc.styles['Normal']
    normal.font.name = '맑은 고딕'
    normal.font.size = Pt(10)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    if demo:
        doc.add_paragraph('예시 모드 · AI 생성 및 법률 사실 검증 미실행')
    doc.add_heading(article.title, level=0)
    for paragraph in article.introduction.split('\n\n'):
        doc.add_paragraph(paragraph)
    for section in article.sections:
        doc.add_heading(section.question, level=1)
        doc.add_paragraph(section.answer)
        for paragraph in section.explanation.split('\n\n'):
            doc.add_paragraph(paragraph)
    doc.add_heading('핵심 체크리스트', level=1)
    for item in article.checklist:
        doc.add_paragraph(item, style='List Bullet')
    doc.add_heading('결론', level=1)
    doc.add_paragraph(article.conclusion)
    doc.add_paragraph(article.cta)
    data = BytesIO()
    doc.save(data)
    return data.getvalue()
